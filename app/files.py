from __future__ import annotations

import hashlib
import io
import json
import mimetypes
from pathlib import Path

from .skillnet.index import BM25Index
from .store import now, uid

MAX_UPLOAD = 24 * 1024 * 1024
TEXT_EXT = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".py",
    ".ts",
    ".tsx",
    ".js",
    ".yaml",
    ".yml",
    ".log",
}


class FileLibrary:
    def __init__(self, store):
        self.store = store
        self.root = store.directory / "files"
        self.root.mkdir(exist_ok=True)

    def save(self, name, content, mime=None, run_id=None, extract=True):
        if len(content) > MAX_UPLOAD:
            raise ValueError("文件超过 24 MB 上限")
        name = Path(name.replace("\\", "/")).name[:180]
        if not name or not content:
            raise ValueError("文件为空")
        ident = uid()
        target = self.root / (ident + Path(name).suffix.lower())
        target.write_bytes(content)
        text, parse_error = "", ""
        if extract:
            try:
                text = self.extract(name, content)[:2_000_000]
            except Exception:
                parse_error = "已保存原文件，暂时无法提取文本"
        self.store.execute(
            "INSERT INTO files VALUES(?,?,?,?,?,?,?,?,?)",
            (
                ident,
                name,
                mime or mimetypes.guess_type(name)[0] or "application/octet-stream",
                len(content),
                hashlib.sha256(content).hexdigest(),
                str(target),
                text,
                run_id,
                now(),
            ),
        )
        return {**self.get(ident), "parse_error": parse_error}

    @staticmethod
    def extract(name, content):
        ext = Path(name).suffix.lower()
        if ext in TEXT_EXT:
            return content.decode("utf-8", errors="replace")
        if ext == ".pdf":
            from pypdf import PdfReader

            return "\n\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(content)).pages
            )
        if ext == ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            return "\n".join(
                [p.text for p in document.paragraphs]
                + [
                    " | ".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
            )
        return ""

    def get(self, ident):
        row = self.store.one("SELECT * FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        return {
            key: value for key, value in row.items() if key not in {"path", "text"}
        } | {"indexed": bool(row["text"])}

    def path(self, ident):
        row = self.store.one("SELECT * FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        path = Path(row["path"]).resolve()
        if not path.is_relative_to(self.root.resolve()) or not path.is_file():
            raise ValueError("文件不可用")
        return path

    def search(self, query, k=5):
        rows = self.store.all(
            "SELECT id,name,text FROM files WHERE text!='' ORDER BY created DESC LIMIT 300"
        )
        chunks = []
        for row in rows:
            text = row["text"][:60000]
            for offset in range(0, len(text), 900):
                chunks.append(
                    {
                        "file_id": row["id"],
                        "filename": row["name"],
                        "offset": offset,
                        "text": text[offset : offset + 1100],
                    }
                )
        if not chunks:
            return []
        index = BM25Index()
        index.fit(
            [str(i) for i in range(len(chunks))], [chunk["text"] for chunk in chunks]
        )
        return [
            {**chunks[int(i)], "score": score}
            for i, score in index.search(query, top_k=k)
        ]

    def read(self, ident):
        row = self.store.one("SELECT name,text FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        return {"id": ident, "name": row["name"], "text": row["text"][:16000]}

    def delete(self, ident):
        target = self.path(ident)
        self.store.execute("DELETE FROM files WHERE id=?", (ident,))
        target.unlink(missing_ok=True)
